"""Text extraction for plain text, Markdown, PDF, and DOCX documents."""

from __future__ import annotations

import logging
import mimetypes
import os
import time
import unicodedata
import zipfile
from pathlib import Path

from smartfs import EXTRACTION_VERSION
from smartfs.config import Settings
from smartfs.models import ExtractedDocument

LOGGER = logging.getLogger("smartfs.extractor")
TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".text"}
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | PDF_EXTENSIONS | DOCX_EXTENSIONS
MAX_EXCERPT_CHARS = 4_000
DEFAULT_TIMEOUT_SECONDS = 120.0


class ExtractionError(Exception):
    """Base class for typed extraction failures."""

    error_kind = "extraction_error"


class UnsupportedFormat(ExtractionError):
    """The file type cannot be indexed by SmartFS."""

    error_kind = "unsupported_format"


class CorruptFile(ExtractionError):
    """The file is malformed or cannot be parsed."""

    error_kind = "corrupt_file"


class EncryptedFile(ExtractionError):
    """The document is encrypted or password protected."""

    error_kind = "encrypted_file"


class EmptyDocument(ExtractionError):
    """The document contains no extractable text."""

    error_kind = "empty_document"


class ExtractionTimeout(ExtractionError):
    """Extraction exceeded its safety deadline."""

    error_kind = "extraction_timeout"


class FileTooLarge(ExtractionError):
    """The file exceeds the configured indexing size limit."""

    error_kind = "file_too_large"


class ExtractorUnavailable(ExtractionError):
    """A required optional extraction library is unavailable."""

    error_kind = "extractor_unavailable"


def mime_for_path(path: str | os.PathLike[str]) -> str | None:
    """Return a best-effort MIME type for a path."""
    mime, _encoding = mimetypes.guess_type(os.fspath(path))
    return mime


def supported_extension(path: str | os.PathLike[str]) -> bool:
    """Return True when the path has an extractable extension."""
    return Path(path).suffix.lower() in SUPPORTED_EXTENSIONS


def _clean_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    output: list[str] = []
    for character in normalized:
        if character in {"\n", "\t"}:
            output.append(character)
        elif unicodedata.category(character).startswith("C"):
            output.append(" ")
        else:
            output.append(character)
    return "".join(output)


def _truncate_text(text: str, limit: int) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit], True


def _deadline_check(started: float, timeout_seconds: float) -> None:
    if time.monotonic() - started > timeout_seconds:
        raise ExtractionTimeout("extraction exceeded its time limit")


def _check_size(path: Path, settings: Settings) -> int:
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise CorruptFile(f"cannot stat file: {exc.strerror}") from exc
    if size > settings.max_file_size_bytes:
        raise FileTooLarge(f"file is {size} bytes; limit is {settings.max_file_size_bytes} bytes")
    return size


def _decode_plain_text(data: bytes) -> tuple[str, str]:
    last_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
            continue
        if text.startswith("\ufeff"):
            text = text.lstrip("\ufeff")
        return text, encoding
    raise CorruptFile("unable to decode text with supported charsets") from last_error


def _extract_text(path: Path, settings: Settings) -> ExtractedDocument:
    size = _check_size(path, settings)
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise CorruptFile(f"cannot read file: {exc.strerror}") from exc
    if b"\x00" in data[:4096]:
        raise CorruptFile("file appears to contain binary data")
    text, encoding = _decode_plain_text(data)
    cleaned = _clean_text(text)
    limited, truncated = _truncate_text(cleaned, settings.max_text_chars)
    if not limited.strip():
        raise EmptyDocument("document contains no readable text")
    return ExtractedDocument(
        text=limited,
        metadata={
            "encoding": encoding,
            "size_bytes": size,
            "mime": mime_for_path(path),
            "extraction_version": EXTRACTION_VERSION,
        },
        page_count=None,
        title=None,
        author=None,
        truncated=truncated,
    )


def _extract_pdf(path: Path, settings: Settings, timeout_seconds: float) -> ExtractedDocument:
    size = _check_size(path, settings)
    try:
        with path.open("rb") as handle:
            header = handle.read(1024)
    except OSError as exc:
        raise CorruptFile(f"cannot read file: {exc.strerror}") from exc
    if b"%PDF-" not in header:
        raise CorruptFile("missing %PDF header")
    try:
        import pymupdf as fitz
    except ImportError as exc:
        raise ExtractorUnavailable("PyMuPDF is not installed") from exc

    started = time.monotonic()
    document = None
    try:
        try:
            document = fitz.open(path)
        except Exception as exc:
            raise CorruptFile(f"PDF could not be opened: {type(exc).__name__}") from exc
        if getattr(document, "needs_pass", False):
            raise EncryptedFile("PDF is password protected")
        page_count = int(document.page_count)
        page_cap = min(page_count, settings.pdf_max_pages)
        parts: list[str] = []
        total_chars = 0
        pages_extracted = 0
        for page_index in range(page_cap):
            _deadline_check(started, timeout_seconds)
            try:
                page_text = document.load_page(page_index).get_text("text")
            except Exception as exc:
                raise CorruptFile(
                    f"PDF page {page_index + 1} could not be read: {type(exc).__name__}"
                ) from exc
            parts.append(page_text)
            total_chars += len(page_text)
            pages_extracted += 1
            if total_chars >= settings.max_text_chars:
                break
        text = _clean_text("\n\n".join(parts))
        limited, text_truncated = _truncate_text(text, settings.max_text_chars)
        if not limited.strip():
            raise EmptyDocument("PDF contains no extractable text")
        raw_metadata = document.metadata or {}
        title = raw_metadata.get("title") or None
        author = raw_metadata.get("author") or None
        return ExtractedDocument(
            text=limited,
            metadata={
                "size_bytes": size,
                "mime": "application/pdf",
                "extraction_version": EXTRACTION_VERSION,
                "pages_extracted": pages_extracted,
                "page_cap": settings.pdf_max_pages,
            },
            page_count=page_count,
            title=str(title)[:500] if title else None,
            author=str(author)[:500] if author else None,
            truncated=bool(text_truncated or page_count > page_cap),
        )
    finally:
        if document is not None:
            document.close()


def _extract_docx(path: Path, settings: Settings, timeout_seconds: float) -> ExtractedDocument:
    size = _check_size(path, settings)
    if not zipfile.is_zipfile(path):
        try:
            with path.open("rb") as handle:
                first = handle.read(8)
        except OSError as exc:
            raise CorruptFile(f"cannot read file: {exc.strerror}") from exc
        if first.startswith(b"\xd0\xcf\x11\xe0"):
            raise EncryptedFile("DOCX appears to be encrypted")
        raise CorruptFile("DOCX is not a valid ZIP package")
    try:
        import docx
    except ImportError as exc:
        raise ExtractorUnavailable("python-docx is not installed") from exc

    started = time.monotonic()
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise CorruptFile(f"DOCX could not be opened: {type(exc).__name__}") from exc

    parts: list[str] = []
    total_chars = 0
    try:
        for paragraph in document.paragraphs:
            _deadline_check(started, timeout_seconds)
            value = paragraph.text.strip()
            if value:
                parts.append(value)
                total_chars += len(value)
            if total_chars >= settings.max_text_chars:
                break
        if total_chars < settings.max_text_chars:
            for table in document.tables:
                for row in table.rows:
                    _deadline_check(started, timeout_seconds)
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        value = " | ".join(cells)
                        parts.append(value)
                        total_chars += len(value)
                    if total_chars >= settings.max_text_chars:
                        break
                if total_chars >= settings.max_text_chars:
                    break
    except ExtractionTimeout:
        raise
    except Exception as exc:
        raise CorruptFile(f"DOCX parsing failed: {type(exc).__name__}") from exc

    text = _clean_text("\n".join(parts))
    limited, truncated = _truncate_text(text, settings.max_text_chars)
    if not limited.strip():
        raise EmptyDocument("DOCX contains no extractable text")
    properties = document.core_properties
    title = properties.title or None
    author = properties.author or None
    return ExtractedDocument(
        text=limited,
        metadata={
            "size_bytes": size,
            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "extraction_version": EXTRACTION_VERSION,
        },
        page_count=None,
        title=str(title)[:500] if title else None,
        author=str(author)[:500] if author else None,
        truncated=truncated,
    )


def extract(
    path: str | os.PathLike[str],
    settings: Settings,
    *,
    timeout_seconds: float = DEFAULT_TIMEOUT_SECONDS,
) -> ExtractedDocument:
    """Extract text and metadata from a supported file.

    Raises a typed :class:`ExtractionError` subclass on failure.
    """
    target = Path(path)
    suffix = target.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return _extract_text(target, settings)
    if suffix in PDF_EXTENSIONS:
        return _extract_pdf(target, settings, timeout_seconds)
    if suffix in DOCX_EXTENSIONS:
        return _extract_docx(target, settings, timeout_seconds)
    raise UnsupportedFormat(f"unsupported extension: {suffix or '(none)'}")


def excerpt(text: str, limit: int = MAX_EXCERPT_CHARS) -> str:
    """Return a bounded text excerpt suitable for indexing."""
    if limit < 1:
        raise ValueError("limit must be positive")
    if len(text) <= limit:
        return text
    cut = text[:limit]
    last_space = cut.rfind(" ")
    if last_space > limit // 2:
        cut = cut[:last_space]
    return cut
